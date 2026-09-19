# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl"]
# ///
"""Turn a typesafe-screening full-results JSON into an Excel workbook.

Usage: uv run results_to_xlsx.py <full_results.json> <output.xlsx>

Sheets: Summary, Include, Maybe, Exclude (+ Error if any). Author, year, journal and
publication type are fetched from PubMed for numeric ids. Never overwrites the output.
"""

import json
import os
import sys
import urllib.parse
import urllib.request
import xml.etree.ElementTree as ET

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

EFETCH_URL = "https://eutils.ncbi.nlm.nih.gov/entrez/eutils/efetch.fcgi"


def fetch_meta(pmids: list[str]) -> dict:
    meta = {}
    for i in range(0, len(pmids), 200):
        params = {"db": "pubmed", "retmode": "xml", "id": ",".join(pmids[i:i + 200])}
        if os.environ.get("NCBI_API_KEY"):
            params["api_key"] = os.environ["NCBI_API_KEY"]
        with urllib.request.urlopen(EFETCH_URL, urllib.parse.urlencode(params).encode(), timeout=60) as resp:
            root = ET.fromstring(resp.read())
        for a in root.iter("PubmedArticle"):
            art = a.find("./MedlineCitation/Article")
            title = art.find("ArticleTitle")
            author = art.find("./AuthorList/Author")
            pubdate = "./Journal/JournalIssue/PubDate/"
            meta[a.findtext("./MedlineCitation/PMID")] = {
                "title": "".join(title.itertext()).strip() if title is not None else "",
                "first": (author.findtext("LastName") or author.findtext("CollectiveName") or "") if author is not None else "",
                "year": art.findtext(pubdate + "Year") or (art.findtext(pubdate + "MedlineDate") or "")[:4],
                "journal": art.findtext("./Journal/ISOAbbreviation") or art.findtext("./Journal/Title") or "",
                "pt": "; ".join(p.text for p in art.iterfind("./PublicationTypeList/PublicationType")),
            }
    return meta


def main() -> None:
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    src, out = sys.argv[1], os.path.expanduser(sys.argv[2])
    if os.path.exists(out):
        sys.exit(f"Refusing to overwrite existing file: {out}")
    with open(os.path.expanduser(src), encoding="utf-8") as f:
        data = json.load(f)
    results = data["results"]
    if isinstance(results, dict):
        sys.exit("This is a slim result. Use the file written by save_full_results_to.")
    inc_texts, exc_texts = data.get("inclusion_criteria", []), data.get("exclusion_criteria", [])
    n_inc = max([len(r.get("inclusion", [])) for r in results] + [len(inc_texts)])
    n_exc = max([len(r.get("exclusion", [])) for r in results] + [len(exc_texts)])
    meta = fetch_meta([r["id"] for r in results if r["id"].isdigit()])

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    rows = [("PubMed query", data.get("pubmed_query", "")), ("Total hits", data.get("total_hits", "")),
            ("Screened", data.get("screened", len(results))), ("Model", data.get("model", "")),
            ("Research question", data.get("research_question", ""))]
    rows += [(f"Inc{i + 1}", t) for i, t in enumerate(inc_texts)] + [(f"Exc{i + 1}", t) for i, t in enumerate(exc_texts)]
    th = data.get("thresholds", {})
    rows.append(("Rule", f"Exc >= {th.get('include')} -> exclude; Match <= {th.get('exclude')} -> exclude; "
                         f"Match >= {th.get('include')} and all Inc >= 0.5 -> include; otherwise maybe"))
    rows += [(k.capitalize(), v) for k, v in data.get("counts", {}).items()]
    for row in rows:
        ws.append(row)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 22, 120
    for c in ws["A"]:
        c.font = Font(bold=True)
    for c in ws["B"]:
        c.alignment = Alignment(wrap_text=True, vertical="top")

    cols = (["PMID", "First author", "Year", "Journal", "Title", "Decision", "Reason", "Match", "Relevance (0-2)",
             "Relevance confidence"] + [f"Inc{i + 1}" for i in range(n_inc)] + [f"Exc{i + 1}" for i in range(n_exc)]
            + ["Publication type", "URL", "Human decision", "Notes"])
    widths = [11, 16, 7, 24, 80, 10, 38, 9, 11, 11] + [8] * (n_inc + n_exc) + [34, 40, 14, 30]
    url_col = cols.index("URL") + 1
    for dec in ("include", "maybe", "exclude", "error"):
        group = [r for r in results if r["decision"] == dec]
        if dec == "error" and not group:
            continue
        ws = wb.create_sheet(dec.capitalize())
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
        for r in group:
            m = meta.get(r["id"], {})
            inc = (r.get("inclusion", []) + [None] * n_inc)[:n_inc]
            exc = (r.get("exclusion", []) + [None] * n_exc)[:n_exc]
            url = f"https://pubmed.ncbi.nlm.nih.gov/{r['id']}/" if r["id"].isdigit() else None
            year = m.get("year", "")
            ws.append([int(r["id"]) if r["id"].isdigit() else r["id"], m.get("first"), int(year) if year.isdigit() else year,
                       m.get("journal"), m.get("title") or r["title"], r["decision"], r["reason"], r.get("match"),
                       r.get("relevance"), r.get("relevance_confidence"), *inc, *exc, m.get("pt"), url, None, None])
            if url:
                cell = ws.cell(ws.max_row, url_col)
                cell.hyperlink, cell.font = url, Font(color="0563C1", underline="single")
            ws.cell(ws.max_row, 5).alignment = Alignment(wrap_text=True, vertical="top")
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
        print(f"{dec}: {len(group)}")
    wb.save(out)
    print(out)


if __name__ == "__main__":
    main()
